"""Deterministic document-to-content pipeline.

Stages (single responsibility each, no overlapping agents):
  1. ingest        - decode a document into raw blocks with metadata
  2. extract       - recover headings, paragraphs, lists and tables (+ OCR fallback)
  3. clean         - normalise whitespace without destroying structure or values
  4. build_ir      - structured intermediate representation of the document
  5. chunk         - semantic chunks carrying page/section metadata

Every stage is a pure function so the pipeline stays inspectable and testable.
"""

import io
import os
import re
from collections import OrderedDict

# ---------------------------------------------------------------- patterns

NUMBER_RE = re.compile(
    r"(?:\$|€|£|USD|EUR|GBP)?\s?\d+(?:[.,]\d+)*(?:\s?(?:%|percent|billion|million|bn|m|k|kg|km|ms|gb|tb|hrs?|hours?|years?|days?|months?))?",
    re.IGNORECASE,
)
DATE_RE = re.compile(
    r"\b(?:19|20)\d{2}\b|\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b|"
    r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)"
    r"(?:\s+\d{1,2},?\s+\d{4})?\b|\bQ[1-4]\s+\d{4}\b",
    re.IGNORECASE,
)
DEFINITION_RE = re.compile(
    r"^(?P<term>[A-Z][A-Za-z0-9 \-]{1,50}?)\s*(?::|—|\bis\b|\bare\b|\brefers to\b|\bmeans\b)\s+(?P<body>.{15,})",
    re.IGNORECASE,
)
LIST_MARKER_RE = re.compile(r"^\s*(?:[-*•‣▪]|\(?\d{1,2}[.)]|[a-z][.)]|[ivx]{1,4}[.)])\s+", re.IGNORECASE)
HEADING_NUMBER_RE = re.compile(r"^\s*\d+(?:\.\d+)*[.)]?\s+\S")


def _clean_inline(text):
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


# ---------------------------------------------------------------- stage 1+2

def _looks_like_heading(line):
    line = line.strip()
    if not line or len(line) > 90:
        return False
    words = line.split()
    if len(words) > 13:
        return False
    if line.endswith((".", ",", ";", ":")):
        return False
    if HEADING_NUMBER_RE.match(line):
        return True
    letters = [c for c in line if c.isalpha()]
    if letters and sum(1 for c in letters if c.isupper()) / len(letters) > 0.7 and len(words) <= 10:
        return True
    if len(words) == 1:
        return line[:1].isupper() and line.isalpha()
    capitalized = sum(1 for w in words if w[:1].isupper())
    return capitalized / len(words) >= 0.7


def _extract_pdf_blocks(file_bytes):
    import pdfplumber
    blocks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        page_count = len(pdf.pages)
        for page_index, page in enumerate(pdf.pages, start=1):
            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []
            page_text = page.extract_text() or ""
            for line in page_text.splitlines():
                line = _clean_inline(line)
                if not line or re.fullmatch(r"\d{1,4}", line):
                    continue
                if _looks_like_heading(line):
                    blocks.append({"type": "heading", "text": line, "page": page_index, "level": 2})
                elif LIST_MARKER_RE.match(line):
                    blocks.append({"type": "list_item", "text": LIST_MARKER_RE.sub("", line), "page": page_index})
                else:
                    blocks.append({"type": "paragraph", "text": line, "page": page_index})
            for table in tables:
                rows = [[_clean_inline(cell or "") for cell in row] for row in table if any(cell for cell in row)]
                if rows:
                    blocks.append({"type": "table", "rows": rows, "page": page_index, "text": ""})
    return blocks, page_count


def _extract_docx_blocks(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    blocks = []
    for paragraph in doc.paragraphs:
        text = _clean_inline(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading") or style.startswith("title"):
            digits = re.findall(r"\d+", style)
            level = int(digits[0]) if digits else 1
            blocks.append({"type": "heading", "text": text, "page": 1, "level": level})
        elif "list" in style or paragraph._p.find(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr"
        ) is not None and paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
            blocks.append({"type": "list_item", "text": LIST_MARKER_RE.sub("", text), "page": 1})
        else:
            blocks.append({"type": "paragraph", "text": text, "page": 1})
    for table in doc.tables:
        rows = [[_clean_inline(cell.text) for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(cell for cell in r)]
        if rows:
            blocks.append({"type": "table", "rows": rows, "page": 1, "text": ""})
    return blocks, 1


def _extract_text_blocks(file_bytes, markdown=False):
    raw = file_bytes.decode("utf-8", errors="replace")
    blocks = []
    pending_table = []
    for line in raw.splitlines():
        stripped = line.strip()
        if markdown and stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append({"type": "heading", "text": stripped.lstrip("#").strip(), "page": 1, "level": level})
            continue
        if stripped.count("|") >= 2:
            pending_table.append([_clean_inline(c) for c in stripped.strip("|").split("|")])
            continue
        if pending_table:
            if not all(re.fullmatch(r"[-: ]+", c or "-") for c in pending_table[0]):
                blocks.append({"type": "table", "rows": pending_table, "page": 1, "text": ""})
            pending_table = []
        if not stripped:
            continue
        if not markdown and _looks_like_heading(stripped):
            blocks.append({"type": "heading", "text": stripped, "page": 1, "level": 2})
        elif LIST_MARKER_RE.match(stripped):
            blocks.append({"type": "list_item", "text": LIST_MARKER_RE.sub("", stripped), "page": 1})
        else:
            blocks.append({"type": "paragraph", "text": stripped, "page": 1})
    if pending_table:
        blocks.append({"type": "table", "rows": pending_table, "page": 1, "text": ""})
    return blocks, 1


def _extract_html_blocks(html_content):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_content, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "aside"]):
        tag.decompose()
    blocks = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "table"]):
        name = element.name
        if name.startswith("h"):
            text = _clean_inline(element.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "heading", "text": text, "page": 1, "level": int(name[1])})
        elif name == "table":
            rows = []
            for tr in element.find_all("tr"):
                cells = [_clean_inline(td.get_text(" ", strip=True)) for td in tr.find_all(["td", "th"])]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append({"type": "table", "rows": rows, "page": 1, "text": ""})
        elif name == "li":
            text = _clean_inline(element.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "list_item", "text": text, "page": 1})
        else:
            text = _clean_inline(element.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "paragraph", "text": text, "page": 1})
    return blocks, 1


def _pdf_needs_ocr(blocks, page_count):
    if page_count == 0:
        return True
    text_chars = sum(len(b.get("text", "")) for b in blocks)
    return (text_chars / max(1, page_count)) < 60


def _try_ocr(file_bytes, page_count):
    """Best-effort OCR. Returns (blocks, engine_name) or (None, reason)."""
    try:
        import pytesseract
        from PIL import Image
    except Exception:
        return None, "OCR engine not installed (pip install pytesseract pillow + a tesseract binary). No text was invented."
    try:
        import pdfplumber
        blocks = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                rendered = page.to_image(resolution=200).original
                text = pytesseract.image_to_string(rendered)
                for line in text.splitlines():
                    line = _clean_inline(line)
                    if line:
                        blocks.append({"type": "paragraph", "text": line, "page": page_index})
        if not blocks:
            return None, "OCR ran but produced no readable text."
        return blocks, "pytesseract"
    except Exception as exc:  # pragma: no cover - depends on local binaries
        return None, f"OCR attempted but failed: {exc}"


# ---------------------------------------------------------------- stage 1-2 API

def extract_blocks(filename, file_bytes):
    """Stage 1+2: return (blocks, metadata). Raises ValueError for unknown types."""
    ext = os.path.splitext(filename or "")[1].lower()
    metadata = {"filename": filename, "extension": ext, "ocr": None, "warnings": []}

    if ext == ".pdf":
        blocks, page_count = _extract_pdf_blocks(file_bytes)
        if _pdf_needs_ocr(blocks, page_count):
            ocr_blocks, note = _try_ocr(file_bytes, page_count)
            if ocr_blocks:
                blocks = ocr_blocks
                metadata["ocr"] = note
            else:
                metadata["warnings"].append(
                    f"Very little text found in a {page_count}-page PDF; it may be scanned. {note}"
                )
        metadata["page_count"] = page_count
    elif ext == ".docx":
        blocks, page_count = _extract_docx_blocks(file_bytes)
        metadata["page_count"] = page_count
    elif ext in (".md", ".markdown"):
        blocks, page_count = _extract_text_blocks(file_bytes, markdown=True)
        metadata["page_count"] = page_count
    elif ext in (".txt", ".text"):
        blocks, page_count = _extract_text_blocks(file_bytes, markdown=False)
        metadata["page_count"] = page_count
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    metadata["block_count"] = len(blocks)
    metadata["has_tables"] = any(b["type"] == "table" for b in blocks)
    metadata["char_count"] = sum(len(b.get("text", "")) for b in blocks) + sum(
        len(cell) for b in blocks if b["type"] == "table" for row in b["rows"] for cell in row
    )
    return blocks, metadata


# ---------------------------------------------------------------- stage 3

def clean_blocks(blocks):
    """Stage 3: normalise text, drop layout noise, keep structure and values."""
    cleaned = []
    previous_text = None
    for block in blocks:
        if block["type"] == "table":
            rows = []
            for row in block["rows"]:
                cells = [_clean_inline(cell) for cell in row]
                if not any(cells):
                    continue
                if all(re.fullmatch(r"[-: ]*", cell or "") for cell in cells):
                    continue
                rows.append(cells)
            if rows:
                cleaned.append({**block, "rows": rows})
            continue
        text = _clean_inline(block.get("text", ""))
        text = re.sub(r"\.{3,}", "…", text)
        text = re.sub(r"\s+([,.;:])", r"\1", text)
        if not text or len(text) < 2:
            continue
        if re.fullmatch(r"(?:page\s*)?\d{1,4}(?:\s*(?:of|/)\s*\d{1,4})?", text, re.IGNORECASE):
            continue
        if previous_text and text == previous_text:
            continue
        previous_text = text
        cleaned.append({**block, "text": text})
    return cleaned


# ---------------------------------------------------------------- stage 4

def build_ir(blocks, metadata):
    """Stage 4: structured intermediate representation."""
    sections = []
    current = {"heading": metadata.get("filename") or "Document", "level": 1, "page": blocks[0]["page"] if blocks else 1,
               "paragraphs": [], "list_items": [], "tables": []}
    title = None

    for block in blocks:
        if block["type"] == "heading":
            if title is None and not current["paragraphs"] and not current["list_items"]:
                title = block["text"]
                current["heading"] = block["text"]
                current["level"] = block.get("level", 1)
                current["page"] = block["page"]
                continue
            sections.append(current)
            current = {"heading": block["text"], "level": block.get("level", 2), "page": block["page"],
                       "paragraphs": [], "list_items": [], "tables": []}
        elif block["type"] == "list_item":
            current["list_items"].append(block["text"])
        elif block["type"] == "table":
            current["tables"].append({"rows": block["rows"], "page": block["page"]})
        else:
            current["paragraphs"].append(block["text"])
    sections.append(current)
    sections = [s for s in sections if s["paragraphs"] or s["list_items"] or s["tables"]]

    if title is None:
        for section in sections:
            if section["paragraphs"]:
                title = section["paragraphs"][0][:120]
                break
    title = title or metadata.get("filename") or "Untitled document"

    full_text = title + "\n\n" + "\n\n".join(
        "\n\n".join([s["heading"]] + s["paragraphs"] + s["list_items"] + [_table_to_text(t["rows"]) for t in s["tables"]])
        for s in sections
    )

    paragraph_pool = [p for s in sections for p in s["paragraphs"]] + [li for s in sections for li in s["list_items"]]

    concepts = _extract_concepts(full_text, title)
    heading_words = set()
    for section in sections:
        heading_words.update(re.split(r"[ \t\-]+", section["heading"]))
    heading_words.update(re.split(r"[ \t\-]+", title or ""))
    entities = _extract_entities(full_text, exclude=[w for w in heading_words if w])
    dates = _unique_in_order(DATE_RE.findall(full_text))[:20]
    numbers = _unique_in_order([n.strip() for n in NUMBER_RE.findall(full_text) if re.search(r"\d", n)])[:40]
    definitions = _extract_definitions(paragraph_pool)
    statements = _extract_statements(paragraph_pool)
    facts = [s for s in statements if NUMBER_RE.search(s)][:12]
    tables = [t for s in sections for t in s["tables"]]

    return {
        "title": title,
        "metadata": metadata,
        "sections": sections,
        "key_concepts": concepts,
        "entities": entities,
        "dates": dates,
        "numbers": numbers,
        "definitions": definitions,
        "important_statements": statements,
        "key_facts": facts,
        "tables": tables,
        "references": metadata.get("source_refs", []),
        "word_count": len(full_text.split()),
        "full_text": full_text,
    }


def _table_to_text(rows):
    return "\n".join(" | ".join(cell for cell in row) for row in rows)


def _unique_in_order(items):
    seen, out = set(), []
    for item in items:
        key = item.lower()
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def _extract_concepts(text, title, limit=14):
    stop = {
        "the", "and", "for", "are", "but", "not", "you", "all", "can", "had", "was", "one", "our",
        "out", "has", "how", "its", "may", "new", "now", "see", "who", "did", "that", "with",
        "this", "from", "they", "have", "will", "would", "there", "their", "what", "about", "which",
        "when", "your", "them", "then", "than", "into", "more", "other", "some", "such", "only",
        "also", "over", "been", "were", "these", "those", "each", "much", "many", "most", "because",
        "between", "through", "during", "before", "after", "while", "where", "being", "does", "using",
        "used", "however", "therefore", "thus", "although", "within", "without", "based", "including",
    }
    words = re.findall(r"[A-Za-z][A-Za-z\-]{3,}", text.lower())
    freq = {}
    for word in words:
        if word in stop:
            continue
        freq[word] = freq.get(word, 0) + 1
    ranked = sorted(freq.items(), key=lambda kv: kv[1], reverse=True)
    concepts = [word for word, count in ranked[:limit] if count > 1]
    return concepts or [word for word, _ in ranked[:limit]]


def _extract_entities(text, exclude=None):
    multi = re.compile(
        r"\b([A-Z]{2,}(?:[ \t]+[A-Z][a-z]+){0,2}|[A-Z][a-z]+(?:[ \t]+[A-Z][a-z]+){1,3})\b"
    )
    camel = re.compile(r"\b([A-Z][a-z]+[A-Z][a-z]+)\b")
    single = re.compile(r"(?<=[a-z,;:()][ \t])([A-Z][a-z]{2,})\b")
    blocklist = {"The Report", "This Report", "The Study", "In Addition", "The Market", "The Company"}
    exclude = {e.lower() for e in (exclude or [])}
    out = []

    def consider(candidate):
        candidate = candidate.strip()
        if not candidate or candidate in blocklist or len(candidate) < 3:
            return
        words = candidate.lower().split()
        if candidate.lower() in exclude or (len(words) > 1 and all(w in exclude for w in words)):
            return
        if candidate not in out:
            out.append(candidate)

    for line in text.splitlines():
        for candidate in multi.findall(line):
            consider(candidate)
        for candidate in camel.findall(line):
            consider(candidate)
        for candidate in single.findall(line):
            consider(candidate)
    return out[:15]


def _extract_definitions(paragraphs):
    definitions = []
    for paragraph in paragraphs:
        match = DEFINITION_RE.match(paragraph)
        if match:
            term = match.group("term").strip()
            body = match.group("body").strip().rstrip(".")
            if 2 <= len(term.split()) <= 6 and 15 <= len(body) <= 320:
                definitions.append({"term": term, "definition": body})
    unique = OrderedDict()
    for item in definitions:
        unique.setdefault(item["term"].lower(), item)
    return list(unique.values())[:12]


def _extract_statements(paragraphs, limit=18):
    statements = []
    for paragraph in paragraphs:
        for sentence in re.split(r"(?<=[.!?])\s+", paragraph):
            sentence = sentence.strip()
            if 40 <= len(sentence) <= 320:
                statements.append(sentence)
    return statements[:limit]


# ---------------------------------------------------------------- stage 5

def chunk_sections(sections, target_words=180, max_words=260):
    """Stage 5: semantic chunks that respect sections and never split a table."""
    chunks = []
    index = 0
    for section in sections:
        buffer, buffer_words, buffer_page = [], 0, section["page"]
        def flush():
            nonlocal index, buffer, buffer_words, buffer_page
            if buffer:
                chunks.append({
                    "chunk_index": index,
                    "section": section["heading"],
                    "level": section["level"],
                    "page": buffer_page,
                    "type": "text",
                    "text": "\n".join(buffer),
                    "word_count": buffer_words,
                })
                index += 1
                buffer, buffer_words = [], 0

        for paragraph in section["paragraphs"] + section["list_items"]:
            words = len(paragraph.split())
            if buffer_words + words > max_words and buffer:
                flush()
                buffer_page = section["page"]
            buffer.append(paragraph)
            buffer_words += words
            if buffer_words >= target_words:
                flush()
        flush()

        for table in section["tables"]:
            chunks.append({
                "chunk_index": index,
                "section": section["heading"],
                "level": section["level"],
                "page": table["page"],
                "type": "table",
                "text": _table_to_text(table["rows"]),
                "word_count": len(_table_to_text(table["rows"]).split()),
            })
            index += 1
    return chunks


def run_pipeline(filename, file_bytes):
    """Run all stages and return a single structured result."""
    blocks, metadata = extract_blocks(filename, file_bytes)
    cleaned = clean_blocks(blocks)
    ir = build_ir(cleaned, metadata)
    ir["source_refs"] = sorted({c["page"] for c in chunk_sections(ir["sections"]) if c.get("page")})
    chunks = chunk_sections(ir["sections"])
    ir["chunks"] = chunks
    ir["stats"] = {
        "blocks": len(cleaned),
        "sections": len(ir["sections"]),
        "chunks": len(chunks),
        "tables": len(ir["tables"]),
        "definitions": len(ir["definitions"]),
        "entities": len(ir["entities"]),
        "dates": len(ir["dates"]),
        "numbers": len(ir["numbers"]),
    }
    return ir

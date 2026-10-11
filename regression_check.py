"""Fast, dependency-light regression checks for TransformIQ.

Run from the project root:
  python regression_check.py --base-url http://127.0.0.1:8003
  python regression_check.py --skip-server
"""
import argparse
import asyncio
import io
import json
import re
import subprocess
import sys
from pathlib import Path

import httpx

from backend.ai_services import content_analyzer, claim_extractor, knowledge_mapper, content_builders, document_pipeline, fact_checker, transformation_engine
from backend.ai_services.provider import AIProvider


def check(name, condition, detail=""):
    if not condition:
        raise AssertionError(f"FAIL {name}: {detail}")
    print(f"PASS {name}")


def make_docx():
    from docx import Document
    doc = Document()
    doc.add_heading("Battery Storage Brief", level=1)
    doc.add_heading("Findings", level=2)
    doc.add_paragraph("Battery storage capacity grew to 12 GWh in 2024.")
    doc.add_paragraph("Costs fell by 28%.", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def direct_checks():
    txt = "A system reduced operating costs by 28% in 2024. The recommendation is to monitor usage."
    ir = document_pipeline.run_pipeline("brief.txt", txt.encode())
    check("TXT extraction", txt in ir["full_text"], ir["full_text"])
    check("numbers and dates", "28%" in ir["numbers"] and "2024" in ir["dates"], str(ir))

    md = "# Battery Brief\n\n## Findings\n\n- Costs fell by 28%\n\n| Metric | Value |\n|---|---|\n| Savings | 28% |"
    md_ir = document_pipeline.run_pipeline("brief.md", md.encode())
    check("Markdown headings", [s["heading"] for s in md_ir["sections"]] == ["Findings"], str(md_ir["sections"]))
    check("Markdown table", len(md_ir["tables"]) == 1, str(md_ir["tables"]))

    docx_ir = document_pipeline.run_pipeline("brief.docx", make_docx())
    check("DOCX heading", docx_ir["title"] == "Battery Storage Brief", docx_ir["title"])
    check("DOCX list/table-safe extraction", "28%" in docx_ir["full_text"], docx_ir["full_text"])

    long = "\n\n".join(f"## Section {i}\nThe documented result is {i}% for the measured cohort." for i in range(40))
    long_ir = document_pipeline.run_pipeline("long.md", long.encode())
    check("semantic chunking", len(long_ir["chunks"]) > 5, str(len(long_ir["chunks"])))
    retrieved = transformation_engine.retrieve_relevant_context(long, [{"claim_text": "result 39% measured cohort"}], {"core_topic": "Section 39", "important_keywords": ["cohort"]}, max_chars=1000)
    check("long-document retrieval", "result is 39%" in retrieved, retrieved[:200])

    dna = {
        "core_topic": "Battery storage economics", "summary": "Costs and capacity changed.",
        "key_arguments": ["Costs fell by 28%"], "key_facts": ["Capacity reached 12 GWh"],
        "statistics": ["Costs fell — 28%", "Capacity — 12 GWh"],
        "important_entities": ["Energy Lab"], "important_keywords": ["battery", "storage"],
        "conclusions": ["Monitor usage"], "definitions": [],
    }
    claims = [
        {"id": 1, "claim_text": "Costs fell by 28% in 2024", "supporting_source_text": "Costs fell by 28% in 2024", "source_location": "Page 1", "importance": "high", "confidence": .9},
    ]
    student = content_builders.build_format_content("Study Notes", "School Student", "Beginner", "Educational", dna, claims)
    executive = content_builders.build_format_content("Executive Summary", "Business Executive", "Intermediate", "Professional", dna, claims)
    check("audience difference", student["content"] != executive["content"], "outputs identical")
    check("format difference", student["content"] != content_builders.build_format_content("Presentation Outline", "School Student", "Beginner", "Educational", dna, claims)["content"], "outputs identical")

    contradiction = asyncio.run(fact_checker.verify_claims("Costs increased by 80%.", claims))
    check("contradiction not verified", contradiction[0]["verification_status"] != "verified", str(contradiction))
    empty_score = asyncio.run(fact_checker.calculate_quality_scores("", "Costs fell by 28%", claims, "General Public", "Summary"))
    check("empty output score is not perfect", empty_score["overall_score"] < 100, str(empty_score))

    provider = AIProvider()
    provider.demo_mode = True
    check("provider demo fallback explicit", provider._demo_response("x"), "empty fallback")

    async def malformed_provider(*args, **kwargs):
        return '{"wrong_shape": true}'

    from backend.ai_services.provider import ai_provider
    original_generate, original_mode = ai_provider.generate, ai_provider.demo_mode
    ai_provider.generate = malformed_provider
    ai_provider.demo_mode = False
    try:
        fallback_dna = asyncio.run(content_analyzer.analyze_content("# Battery\n\nBattery costs fell by 28% in 2024."))
        fallback_claims = asyncio.run(claim_extractor.extract_claims("Battery costs fell by 28% in 2024."))
        fallback_nodes = asyncio.run(knowledge_mapper.build_knowledge_map("# Battery\n\nBattery costs fell by 28% in 2024."))
        check("malformed DNA fallback", bool(fallback_dna.get("core_topic")), str(fallback_dna))
        check("malformed claims fallback", isinstance(fallback_claims, list), str(fallback_claims))
        check("malformed map fallback", isinstance(fallback_nodes, list), str(fallback_nodes))
    finally:
        ai_provider.generate, ai_provider.demo_mode = original_generate, original_mode


def api_checks(base_url):
    c = httpx.Client(base_url=base_url, timeout=60)
    health = c.get("/api/health")
    check("health endpoint", health.status_code == 200 and health.json().get("status") == "ok", health.text)

    empty = c.post("/api/sources/upload", data={"text": "   "})
    check("empty text rejected", empty.status_code == 400, empty.text)

    upload = c.post("/api/sources/upload", data={"text": "# Supply Chain\n\nCosts fell by 28% in 2024."})
    check("direct text upload", upload.status_code == 200, upload.text)
    source_id = upload.json()["source_id"]
    check("structure returned", upload.json().get("structure", {}).get("stats", {}).get("blocks", 0) > 0, upload.text)

    analysis = c.post("/api/sources/analyze", json={"source_id": source_id})
    check("analysis response", analysis.status_code == 200 and analysis.json()["claims_count"] > 0, analysis.text)

    output = c.post("/api/transform", json={
        "source_id": source_id, "output_format": "Executive Summary", "output_formats": ["Executive Summary"],
        "audience": "Business Executive", "knowledge_level": "Intermediate", "tone": "Professional", "length": "short",
    })
    check("transformation response", output.status_code == 200, output.text)
    item = output.json()["outputs"][0]
    check("validation returned", isinstance(item.get("validation"), dict), str(item))
    check("audience returned", item.get("target_audience") == "Business Executive", str(item))
    check("plan returned", isinstance(item.get("transformation_plan"), dict), str(item))

    chat = c.post("/api/chat", json={"message": "Add the key statistics", "transformation_id": item["transformation_id"], "current_output": item["content"]})
    check("chat response", chat.status_code == 200 and chat.json().get("modified_content"), chat.text)

    exported = c.post("/api/export", json={"content": "hello", "format": "json", "filename": "qa", "metadata": {"source_id": source_id}})
    check("JSON export", exported.status_code == 200 and exported.json()["filename"].endswith(".json"), exported.text)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8003")
    parser.add_argument("--skip-server", action="store_true")
    args = parser.parse_args()
    direct_checks()
    if not args.skip_server:
        api_checks(args.base_url)
    print("ALL REGRESSION CHECKS PASSED")


if __name__ == "__main__":
    main()
